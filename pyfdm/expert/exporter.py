# Copyright (c) 2026 Jakub Więckowski

import csv
import json
import numpy as np
from pathlib import Path

__all__ = ['FuzzyExporter']


class FuzzyExporter:
    """
    Export TFN decision matrices collected by ExpertCollector.

    Supports CSV, JSON, and Excel (.xlsx) output.
    Excel export requires the optional dependency openpyxl:
    ``pip install pyfdm[excel]`` or ``pip install openpyxl``.

    Parameters
    ----------
        collector : ExpertCollector
            A populated ExpertCollector instance.

        alternative_names : list of str, optional
            Labels for alternatives (A1, A2, … by default).

        criterion_names : list of str, optional
            Labels for criteria (C1, C2, … by default).

    Examples
    --------
        >>> exporter = FuzzyExporter(collector)
        >>> exporter.to_csv('results.csv')
        >>> exporter.to_json('results.json')
        >>> exporter.to_excel('results.xlsx')
    """

    def __init__(
        self, 
        collector,
        alternative_names: list | None = None,
        criterion_names: list | None = None
    ):
        self.collector = collector
        m, n = collector.n_alternatives, collector.n_criteria
        self.alt_names = alternative_names or [f'A{i+1}' for i in range(m)]
        self.crit_names = criterion_names or [f'C{j+1}' for j in range(n)]

    #  CSV                                                                 
    def to_csv(self, path: str | Path) -> None:
        """
        Export all expert matrices to a single CSV file.

        Layout: one block per expert, separated by a blank row.
        Each TFN is expanded into three columns: <CritName>_l, _m, _u.

        Parameters
        ----------
            path : str or Path
        """
        path = Path(path)
        n = self.collector.n_criteria

        # Build header
        header = ['Expert', 'Alternative'] + [
            f'{c}_{suffix}'
            for c in self.crit_names
            for suffix in ('l', 'm', 'u')
        ]

        with open(path, 'w', newline='', encoding='utf-8') as f:
            writer = csv.writer(f)
            writer.writerow(header)

            for eid in self.collector.expert_ids:
                matrix = self.collector.get_tfn_matrix(eid)
                for i, alt in enumerate(self.alt_names):
                    row = [f'E{eid}', alt]
                    for j in range(n):
                        row += matrix[i, j].tolist()
                    writer.writerow(row)
                writer.writerow([])  # blank separator

    #  JSON                                                                
    def to_json(self, path: str | Path, indent: int = 2) -> None:
        """
        Export all expert matrices to a JSON file.

        Structure::

            {
              "metadata": { "n_alternatives": ..., "n_criteria": ..., "n_experts": ... },
              "alternatives": [...],
              "criteria": [...],
              "experts": {
                "0": {
                  "A1": { "C1": [l, m, u], "C2": [...], ... },
                  ...
                },
                ...
              }
            }

        Parameters
        ----------
            path : str or Path
            indent : int, default 2
        """
        path = Path(path)
        data = {
            'metadata': {
                'n_alternatives': self.collector.n_alternatives,
                'n_criteria': self.collector.n_criteria,
                'n_experts': self.collector.n_experts,
            },
            'alternatives': self.alt_names,
            'criteria': self.crit_names,
            'experts': {}
        }

        for eid in self.collector.expert_ids:
            matrix = self.collector.get_tfn_matrix(eid)
            expert_data = {}
            for i, alt in enumerate(self.alt_names):
                expert_data[alt] = {
                    crit: matrix[i, j].tolist()
                    for j, crit in enumerate(self.crit_names)
                }
            data['experts'][str(eid)] = expert_data

        with open(path, 'w', encoding='utf-8') as f:
            json.dump(data, f, indent=indent, ensure_ascii=False)

    #  Excel                                                               
    def to_excel(self, path: str | Path) -> None:
        """
        Export all expert matrices to an Excel workbook (.xlsx).

        Creates one worksheet per expert plus a "Summary" sheet showing
        the element-wise mean TFN across all experts.

        Requires ``openpyxl``:  pip install pyfdm[excel]

        Parameters
        ----------
            path : str or Path

        Raises
        ------
            ImportError – if openpyxl is not installed.
        """
        try:
            import openpyxl
            from openpyxl.styles import (
                PatternFill, Font, Alignment, Border, Side
            )
        except ImportError as e:
            raise ImportError(
                'Excel export requires openpyxl. '
                'Install it with: pip install pyfdm[excel]'
            ) from e

        path = Path(path)
        wb = openpyxl.Workbook()
        wb.remove(wb.active)   # remove default sheet

        # Colour palette
        BLUE  = 'FF1F4E79'
        LBLUE = 'FFD6E4F0'
        GREEN = 'FFE2EFDA'
        WHITE = 'FFFFFFFF'
        GRAY  = 'FFF2F2F2'

        thin = Side(style='thin', color='FFCCCCCC')
        def borders():
            return Border(left=thin, right=thin, top=thin, bottom=thin)

        def header_fill():
            return PatternFill('solid', fgColor=BLUE)

        def subheader_fill():
            return PatternFill('solid', fgColor=LBLUE)

        def header_font():
            return Font(bold=True, color='FFFFFFFF', name='Arial', size=10)

        def normal_font():
            return Font(name='Arial', size=10)

        n = self.collector.n_criteria

        def _write_matrix_sheet(ws, matrix, sheet_title):
            """Write one TFN matrix to a worksheet."""
            # Title row
            ws.merge_cells(
                start_row=1, start_column=1,
                end_row=1, end_column=1 + n * 3
            )
            title_cell = ws.cell(row=1, column=1, value=sheet_title)
            title_cell.fill = PatternFill('solid', fgColor=BLUE)
            title_cell.font = Font(bold=True, color='FFFFFFFF', name='Arial', size=12)
            title_cell.alignment = Alignment(horizontal='center', vertical='center')

            # Criterion header (spanning 3 columns each)
            ws.cell(row=2, column=1, value='Alternative').fill = subheader_fill()
            ws.cell(row=2, column=1).font = header_font()
            ws.cell(row=2, column=1).border = borders()

            for j, crit in enumerate(self.crit_names):
                col = 2 + j * 3
                ws.merge_cells(
                    start_row=2, start_column=col,
                    end_row=2, end_column=col + 2
                )
                cell = ws.cell(row=2, column=col, value=crit)
                cell.fill = subheader_fill()
                cell.font = header_font()
                cell.alignment = Alignment(horizontal='center')
                cell.border = borders()

            # Sub-header: l, m, u
            for j in range(n):
                for k, suffix in enumerate(['l', 'm', 'u']):
                    col = 2 + j * 3 + k
                    c = ws.cell(row=3, column=col, value=suffix)
                    c.fill = PatternFill('solid', fgColor='FFBDD7EE')
                    c.font = Font(bold=True, name='Arial', size=9)
                    c.alignment = Alignment(horizontal='center')
                    c.border = borders()

            ws.cell(row=3, column=1, value='').border = borders()

            # Data rows
            for i, alt in enumerate(self.alt_names):
                row = 4 + i
                fill = PatternFill('solid', fgColor=GRAY if i % 2 == 1 else WHITE)

                alt_cell = ws.cell(row=row, column=1, value=alt)
                alt_cell.font = Font(bold=True, name='Arial', size=10)
                alt_cell.fill = fill
                alt_cell.border = borders()

                for j in range(n):
                    for k in range(3):
                        col = 2 + j * 3 + k
                        val = round(float(matrix[i, j, k]), 4)
                        c = ws.cell(row=row, column=col, value=val)
                        c.fill = fill
                        c.font = normal_font()
                        c.number_format = '0.0000'
                        c.alignment = Alignment(horizontal='center')
                        c.border = borders()

            # Column widths
            ws.column_dimensions['A'].width = 14
            for j in range(n):
                for k in range(3):
                    col_letter = openpyxl.utils.get_column_letter(2 + j * 3 + k)
                    ws.column_dimensions[col_letter].width = 8

        # Expert sheets
        all_matrices = []
        for eid in self.collector.expert_ids:
            matrix = self.collector.get_tfn_matrix(eid)
            all_matrices.append(matrix)
            ws = wb.create_sheet(title=f'Expert {eid}')
            _write_matrix_sheet(ws, matrix, f'Expert {eid} — TFN Decision Matrix')

        # Summary sheet
        if all_matrices:
            mean_matrix = np.mean(all_matrices, axis=0)
            ws_sum = wb.create_sheet(title='Summary (Mean)', index=0)
            _write_matrix_sheet(
                ws_sum, mean_matrix,
                f'Summary — Arithmetic Mean of {len(all_matrices)} Expert(s)'
            )

        wb.save(path)
